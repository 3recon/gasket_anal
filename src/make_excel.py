"""Build the Excel deliverable: data (raw + modelling), sources, statistics,
hyperparameter-trial log, best models, predictions and BO recommendations."""
import json
import os
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gp_core import FILLER_COLS, TARGETS, add_derived, target_rows

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.environ.get("OUT_DIR", os.path.join(ROOT, "outputs"))
XLSX = os.environ.get("XLSX", os.path.join(ROOT, "가스켓_배합-물성_데이터_및_GP모델.xlsx"))

raw = pd.read_csv(os.path.join(ROOT, "data", "gasket_dataset_raw.csv"))
mdl = add_derived(pd.read_csv(os.path.join(ROOT, "data", "gasket_dataset_clean.csv")))
res = json.load(open(os.path.join(OUT, "results.json")))
trials = pd.read_csv(os.path.join(OUT, "trials.csv"))

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True, color="1F3A5F")
INPUT_FONT = Font(name=FONT, size=10, color="0000FF")   # hard-coded source data
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(bottom=thin)

wb = Workbook()
wb.calculation = CalcProperties(fullCalcOnLoad=True)


def write_table(ws, df, start_row=1, input_cols=(), widths=None, numfmt=None):
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(start_row, j, c)
        cell.font, cell.fill = HDR_FONT, HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, float) and np.isnan(v):
                v = None
            if isinstance(v, (np.floating,)):
                v = float(v)
            if isinstance(v, (np.integer,)):
                v = int(v)
            if isinstance(v, (np.bool_,)):
                v = bool(v)
            cell = ws.cell(i, j, v)
            cell.font = INPUT_FONT if df.columns[j - 1] in input_cols else BODY
            if numfmt and df.columns[j - 1] in numfmt:
                cell.number_format = numfmt[df.columns[j - 1]]
    ws.freeze_panes = ws.cell(start_row + 1, 1)
    for j, c in enumerate(df.columns, 1):
        w = (widths or {}).get(c, min(max(10, len(str(c)) + 2), 28))
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[start_row].height = 30
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(df.columns))}{start_row + len(df)}"


# ------------------------------------------------------------------ README
ws = wb.active
ws.title = "README"
lines = [
    ("반도체 장비용 불소고무 가스켓: 배합비-물성 데이터셋 및 가우시안 프로세스 모델", TITLE),
    ("", BODY),
    ("목적", BOLD),
    ("공개 특허의 실시예/비교예 표에서 원재료(고무) 및 부재료(필러, 가교제, 가교조제, 수산제 등)의 배합비(phr)와 "
     "물성(경도, 인장강도, 신율, 100% 모듈러스, 압축영구줄음률)을 추출하고, 가우시안 프로세스(GP) 회귀 + "
     "베이지안 최적화(BO)로 물성 예측 모델을 구축한 결과를 담은 파일입니다.", BODY),
    ("", BODY),
    ("시트 구성", BOLD),
    ("Sources : 데이터 출처 특허 목록 (특허번호, 출원인, 제목, URL, 수집 행수/모델 사용 행수 = COUNTIF 수식)", BODY),
    ("Raw_Data : 특허에서 추출한 전체 원자료 (표 값을 그대로 전사, 단위만 MPa로 환산)", BODY),
    ("Model_Data : 정제 규칙을 통과해 모델링에 사용한 데이터 (파생변수 total_filler_phr 은 SUM 수식)", BODY),
    ("Stats : 물성별 데이터 수·평균·표준편차·최소·최대 (Model_Data 참조 수식)", BODY),
    ("HP_Trials : 하이퍼파라미터 실험 전체 기록 (기준모델, 커널 스크리닝, BO 3회 반복, 강건성 재평가)", BODY),
    ("Best_Models : 물성별 최종 선택 하이퍼파라미터와 교차검증/테스트/출처별 검증 성능", BODY),
    ("Predictions : 개발셋 OOF 예측 및 테스트셋 예측 (예측 표준편차, 95% 구간, 잔차 수식 포함)", BODY),
    ("BO_Recommend : 학습된 GP + 제약 기대개선(cEI)으로 제안한 차기 실험 배합 후보 (실험 검증 전 예측값)", BODY),
    ("", BODY),
    ("데이터 무결성 원칙", BOLD),
    ("모든 수치는 특허 원문 표/본문에서 전사했습니다. 추정·보간·합성 데이터는 포함하지 않았습니다. "
     "보고되지 않은 값은 빈 칸입니다. 파란 글씨 = 특허에서 가져온 입력값, 검정 = 수식/계산값.", BODY),
    ("단위 환산: kgf/cm² × 0.0980665 = MPa, psi × 0.00689476 = MPa. phr = 고무 100 중량부 기준 부재료 중량부.", BODY),
    ("", BODY),
    ("정제 규칙 (Raw_Data → Model_Data)", BOLD),
    ("1) 고무 종류가 FFKM/FKM/FEPM 가 아닌 블렌드(실리콘/EPDM 혼합 등) 제외", BODY),
    ("2) Shore A 가 아닌 경도(마이크로 경도계 Shore M 등)는 경도 열에서 제외", BODY),
    ("3) 100% 모듈러스가 보고되었는데 신율이 100% 미만인 경우 오기로 판단해 신율 제외", BODY),
    ("4) 서로 다른 특허(패밀리)에 동일 배합·동일 물성으로 중복 수록된 행 제거", BODY),
    ("5) 압축영구줄음률은 시험 온도가 명시되고 시험시간 60~100 h 인 행만 사용하며, 시험온도(cs_temp_C)를 입력변수로 추가", BODY),
    ("", BODY),
    ("요약 수치", BOLD),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(i, 1, t)
    c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")
r0 = len(lines) + 1
summary = [
    ("수집 원자료 행 수", "=COUNTA(Raw_Data!A:A)-1"),
    ("수집 출처(특허) 수", "=COUNTA(Sources!A:A)-2"),   # minus header and the total row
    ("모델링 사용 행 수", "=COUNTA(Model_Data!A:A)-1"),
]
for k, (lab, f) in enumerate(summary):
    ws.cell(r0 + k, 1, lab).font = BODY
    c = ws.cell(r0 + k, 2, f)
    c.font = BOLD
ws.column_dimensions["A"].width = 120
ws.column_dimensions["B"].width = 14

# ------------------------------------------------------------------ Sources
src = (raw.groupby("source_id")
       .agg(assignee=("assignee", "first"), title=("title", "first"), url=("url", "first"))
       .reset_index())
ws = wb.create_sheet("Sources")
write_table(ws, src.assign(rows_collected=None, rows_in_model=None),
            widths={"source_id": 18, "assignee": 30, "title": 60, "url": 48,
                    "rows_collected": 14, "rows_in_model": 14})
for i in range(2, len(src) + 2):
    ws.cell(i, 5, f"=COUNTIF(Raw_Data!$A:$A,A{i})").font = BODY
    ws.cell(i, 6, f"=COUNTIF(Model_Data!$B:$B,A{i})").font = BODY
    u = ws.cell(i, 4)
    if u.value:
        u.hyperlink = u.value
        u.font = Font(name=FONT, size=10, color="0563C1", underline="single")
n = len(src) + 2
ws.cell(n, 1, "합계").font = BOLD
ws.cell(n, 5, f"=SUM(E2:E{n - 1})").font = BOLD
ws.cell(n, 6, f"=SUM(F2:F{n - 1})").font = BOLD

# ------------------------------------------------------------------ Raw_Data
schema_cols = [c for c in raw.columns if c != "batch"]
ws = wb.create_sheet("Raw_Data")
num_cols = [c for c in schema_cols if c.endswith(("_phr", "_MPa", "_pct", "_shoreA"))]
write_table(ws, raw[schema_cols], input_cols=set(schema_cols),
            widths={"title": 40, "url": 30, "other_ingredients_note": 40, "extraction_notes": 60,
                    "cs_condition": 30, "plasma_condition": 30, "units_original": 30})

# ------------------------------------------------------------------ Model_Data
ws = wb.create_sheet("Model_Data")
mcols = (["row_id", "source_id", "assignee", "example_label", "polymer_type", "polymer_grade",
          "cure_system", "polymer_phr"] +
         ["carbon_black_phr", "silica_phr", "alumina_phr", "barium_sulfate_phr", "titanium_oxide_phr",
          "fluororesin_phr", "other_inorganic_filler_phr", "organic_additive_phr", "coagent_phr",
          "peroxide_phr", "nitrile_curative_phr", "other_curative_phr", "acid_acceptor_phr",
          "processing_aid_phr"] +
         ["hardness_shoreA", "tensile_MPa", "elongation_pct", "m100_MPa", "compression_set_pct",
          "cs_condition", "cs_temp_C"])
md = mdl[mcols].copy()
md["total_filler_phr"] = None
for t in TARGETS:
    md[f"used_{t}"] = mdl.index.isin(target_rows(mdl, t).index).astype(int)
write_table(ws, md, input_cols=set(mcols) - {"row_id", "cs_temp_C"},
            widths={"polymer_grade": 26, "cs_condition": 30, "example_label": 22, "assignee": 22})
col = {c: get_column_letter(j) for j, c in enumerate(md.columns, 1)}
fl = [col[c] for c in FILLER_COLS]
for i in range(2, len(md) + 2):
    ws[f"{col['total_filler_phr']}{i}"] = "=" + "+".join(f"{c}{i}" for c in fl)
    ws[f"{col['total_filler_phr']}{i}"].font = BODY
MD_N = len(md) + 1

# ------------------------------------------------------------------ Stats
ws = wb.create_sheet("Stats")
hdr = ["물성", "열", "n (전체)", "평균", "표준편차", "최소", "최대", "n (모델 사용)"]
for j, h in enumerate(hdr, 1):
    c = ws.cell(1, j, h)
    c.font, c.fill = HDR_FONT, HDR_FILL
names = {"hardness_shoreA": "경도 (Shore A)", "tensile_MPa": "인장강도 (MPa)", "elongation_pct": "신율 (%)",
         "m100_MPa": "100% 모듈러스 (MPa)", "compression_set_pct": "압축영구줄음률 (%)"}
for i, t in enumerate(TARGETS, 2):
    L = col[t]
    rng = f"Model_Data!${L}$2:${L}${MD_N}"
    U = col[f"used_{t}"]
    ws.cell(i, 1, names[t]).font = BODY
    ws.cell(i, 2, t).font = BODY
    ws.cell(i, 3, f"=COUNT({rng})").font = BODY
    for j, fn in zip(range(4, 8), ["AVERAGE", "STDEV", "MIN", "MAX"]):
        c = ws.cell(i, j, f"={fn}({rng})")
        c.font, c.number_format = BODY, "0.00"
    ws.cell(i, 8, f"=SUM(Model_Data!${U}$2:${U}${MD_N})").font = BODY
for j, w in enumerate([20, 22, 10, 10, 10, 10, 10, 14], 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.cell(8, 1, "주: 압축영구줄음률의 '모델 사용' 수는 시험온도 명시 및 60~100 h 조건을 만족하는 행만 집계").font = BODY

# ------------------------------------------------------------------ HP_Trials
ws = wb.create_sheet("HP_Trials")
tcols = ["target", "experiment", "seed", "iter", "model", "kernel", "ard", "linear", "x_transform",
         "y_transform", "feature_set", "ls_init", "noise_init", "noise_floor", "n_restarts",
         "R2", "R2_std", "RMSE", "MAE", "cov95", "NLPD", "sec"]
tt = trials[[c for c in tcols if c in trials.columns]].copy()
tt.insert(0, "trial_no", range(1, len(tt) + 1))
write_table(ws, tt, numfmt={c: "0.000" for c in ["R2", "R2_std", "RMSE", "MAE", "cov95", "NLPD",
                                                   "ls_init", "noise_init", "noise_floor"]} | {"sec": "0.0"})

# ------------------------------------------------------------------ Best_Models
rows = []
for t, R in res.items():
    c = R["best_config"]
    rows.append({
        "target": t, "n_total": R["n_total"], "n_dev": R["n_dev"], "n_test": R["n_test"],
        "n_sources": R["n_sources"], "kernel": c["kernel"], "ARD": c["ard"] and c["kernel"] != "rq",
        "linear_term": c["linear"], "x_transform": c["x_transform"], "y_transform": c["y_transform"],
        "feature_set": c["feature_set"], "ls_init(log10)": c["ls_init"], "noise_init(log10)": c["noise_init"],
        "noise_floor(log10)": c["noise_floor"], "n_restarts(final)": c["n_restarts"],
        "CV_R2(5x5)": R["best_cv"]["R2"], "CV_R2_std": R["best_cv"]["R2_std"], "CV_RMSE": R["best_cv"]["RMSE"],
        "Test_R2": R["test"]["R2"], "Test_RMSE": R["test"]["RMSE"], "Test_MAE": R["test"]["MAE"],
        "Test_95PI_coverage": R["test"]["cov95"], "LeaveSourceOut_R2": R["leave_source_out"]["R2"],
        "RF_Test_R2": R["test_baselines"]["RandomForest"]["R2"], "Ridge_Test_R2": R["test_baselines"]["Ridge"]["R2"],
        "GPdefault_Test_R2": R["test_baselines"]["GP_sklearn_default"]["R2"],
        "final_kernel (fitted on all data)": R["final_kernel"], "log_marginal_likelihood": R["final_lml"],
    })
bm = pd.DataFrame(rows)
ws = wb.create_sheet("Best_Models")
write_table(ws, bm, widths={"final_kernel (fitted on all data)": 90},
            numfmt={c: "0.000" for c in bm.columns if any(k in c for k in ("R2", "RMSE", "MAE", "coverage",
                                                                             "log10", "likelihood"))})

# ------------------------------------------------------------------ Predictions
ws = wb.create_sheet("Predictions")
P = []
for t in res:
    p = pd.read_csv(os.path.join(OUT, f"pred_{t}.csv"))
    p.insert(0, "target", t)
    P.append(p)
P = pd.concat(P, ignore_index=True)
P["lower95"], P["upper95"], P["residual"], P["inside95"] = None, None, None, None
write_table(ws, P, numfmt={c: "0.00" for c in ["y_true", "y_pred", "y_std", "lower95", "upper95", "residual"]})
pc = {c: get_column_letter(j) for j, c in enumerate(P.columns, 1)}
for i in range(2, len(P) + 2):
    yp, ys, yt = f"{pc['y_pred']}{i}", f"{pc['y_std']}{i}", f"{pc['y_true']}{i}"
    ws[f"{pc['lower95']}{i}"] = f"={yp}-1.96*{ys}"
    ws[f"{pc['upper95']}{i}"] = f"={yp}+1.96*{ys}"
    ws[f"{pc['residual']}{i}"] = f"={yp}-{yt}"
    ws[f"{pc['inside95']}{i}"] = f"=IF(AND({yt}>={pc['lower95']}{i},{yt}<={pc['upper95']}{i}),1,0)"
    for k in ("lower95", "upper95", "residual", "inside95"):
        ws[f"{pc[k]}{i}"].font = BODY
        if k != "inside95":
            ws[f"{pc[k]}{i}"].number_format = "0.00"

# ------------------------------------------------------------------ BO_Recommend
rec_path = os.path.join(OUT, "recommendations.csv")
if os.path.exists(rec_path):
    rec = pd.read_csv(rec_path)
    ws = wb.create_sheet("BO_Recommend")
    ws.cell(1, 1, "주의: 아래는 GP 모델 예측값(평균±표준편차)에 기반한 '다음 실험 후보'이며 실측 데이터가 아닙니다. "
                  "실제 배합·가황 후 시험으로 검증이 필요합니다.").font = Font(name=FONT, size=10, bold=True,
                                                                   color="C00000")
    write_table(ws, rec, start_row=3, numfmt={c: "0.00" for c in rec.columns if rec[c].dtype.kind == "f"},
                widths={"nearest_example": 34})

wb.save(XLSX)
print("saved", XLSX)
